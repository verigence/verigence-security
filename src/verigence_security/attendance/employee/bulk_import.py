from __future__ import annotations

from datetime import UTC, date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from uuid import UUID, uuid4

from openpyxl import Workbook, load_workbook
from sqlalchemy import Connection, text

from verigence_security.attendance.employee.errors import (
    AttendanceNotFoundError,
    AttendanceRuleError,
)
from verigence_security.attendance.employee.repository import create_employee

_HEADERS = [
    "EmployeeCode",
    "SecurityUserId",
    "DisplayName",
    "PrimaryEmail",
    "Mobile",
    "JoiningDate",
    "TLUserId",
    "PMOUserId",
    "ProjectTenantId",
    "WorkLocationId",
    "BankAccountMasked",
    "PANMasked",
    "AadhaarMasked",
    "BasicSalary",
    "HRA",
    "Allowances",
    "OtherEarnings",
    "FixedDeductions",
    "SalaryEffectiveFrom",
]


def build_employee_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Employees"
    sheet.append(_HEADERS)
    sheet.append(
        [
            "EMP001",
            "00000000-0000-4000-8000-000000000001",
            "Sample Employee",
            "employee@example.com",
            "+919876543210",
            datetime.now(UTC).date(),
            "",
            "",
            "",
            "",
            "XXXX1234",
            "ABCDE1234F",
            "XXXX-XXXX-1234",
            30000,
            12000,
            5000,
            0,
            2000,
            datetime.now(UTC).date(),
        ]
    )
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column) + 2
        sheet.column_dimensions[column[0].column_letter].width = min(max(width, 12), 28)
    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


def _cell_string(value: object) -> str | None:
    if value is None:
        return None
    normalized = str(value).strip()
    return normalized or None


def _uuid(value: object, *, field_name: str, required: bool = False) -> str | None:
    raw = _cell_string(value)
    if raw is None:
        if required:
            raise ValueError(f"{field_name} is required")
        return None
    try:
        return str(UUID(raw))
    except ValueError as exc:
        raise ValueError(f"{field_name} must be a UUID") from exc


def _date(value: object, *, field_name: str, required: bool = False) -> date | None:
    if value is None or value == "":
        if required:
            raise ValueError(f"{field_name} is required")
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value).strip()
    try:
        return date.fromisoformat(raw)
    except ValueError as exc:
        raise ValueError(f"{field_name} must be YYYY-MM-DD") from exc


def _money(value: object, *, field_name: str) -> Decimal:
    if value is None or value == "":
        return Decimal(0)
    try:
        amount = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{field_name} must be numeric") from exc
    if amount < 0:
        raise ValueError(f"{field_name} cannot be negative")
    return amount.quantize(Decimal("0.01"))


def _normalize_row(values: list[object], row_number: int) -> tuple[dict[str, object], list[str]]:
    row = dict(zip(_HEADERS, values, strict=False))
    messages: list[str] = []
    payload: dict[str, object] = {}
    try:
        code = _cell_string(row.get("EmployeeCode"))
        if not code:
            raise ValueError("EmployeeCode is required")
        payload["employee_code"] = code
        payload["security_user_id"] = _uuid(
            row.get("SecurityUserId"),
            field_name="SecurityUserId",
            required=True,
        )
        display_name = _cell_string(row.get("DisplayName"))
        if not display_name:
            raise ValueError("DisplayName is required")
        payload["display_name"] = display_name
        payload["primary_email"] = _cell_string(row.get("PrimaryEmail"))
        payload["mobile"] = _cell_string(row.get("Mobile"))
        payload["joining_date"] = _date(
            row.get("JoiningDate"),
            field_name="JoiningDate",
            required=True,
        )
        payload["tl_user_id"] = _uuid(row.get("TLUserId"), field_name="TLUserId")
        payload["pmo_user_id"] = _uuid(row.get("PMOUserId"), field_name="PMOUserId")
        payload["project_tenant_id"] = _uuid(
            row.get("ProjectTenantId"),
            field_name="ProjectTenantId",
        )
        payload["work_location_id"] = _uuid(
            row.get("WorkLocationId"),
            field_name="WorkLocationId",
        )
        payload["bank_account_masked"] = _cell_string(row.get("BankAccountMasked"))
        payload["pan_masked"] = _cell_string(row.get("PANMasked"))
        payload["aadhaar_masked"] = _cell_string(row.get("AadhaarMasked"))
        payload["basic_salary"] = str(_money(row.get("BasicSalary"), field_name="BasicSalary"))
        payload["hra"] = str(_money(row.get("HRA"), field_name="HRA"))
        payload["allowances"] = str(_money(row.get("Allowances"), field_name="Allowances"))
        payload["other_earnings"] = str(
            _money(row.get("OtherEarnings"), field_name="OtherEarnings")
        )
        payload["fixed_deductions"] = str(
            _money(row.get("FixedDeductions"), field_name="FixedDeductions")
        )
        salary_effective = _date(
            row.get("SalaryEffectiveFrom"),
            field_name="SalaryEffectiveFrom",
        )
        payload["salary_effective_from"] = str(
            salary_effective or payload["joining_date"]
        )
    except ValueError as exc:
        messages.append(str(exc))
    payload["row_number"] = row_number
    return payload, messages


def _existing_employee(
    connection: Connection,
    *,
    employee_code: str,
    security_user_id: str,
) -> dict[str, Any] | None:
    rows = list(
        connection.execute(
            text(
                """
                SELECT employee_id,employee_code,security_user_id
                FROM verigence_attendance.employees
                WHERE employee_code=:employee_code
                   OR security_user_id=CAST(:security_user_id AS uuid)
                """
            ),
            {
                "employee_code": employee_code,
                "security_user_id": security_user_id,
            },
        ).mappings()
    )
    if not rows:
        return None
    if len(rows) > 1:
        raise AttendanceRuleError(
            "EMPLOYEE_IMPORT_CONFLICT",
            "Employee code and Security user resolve to different existing employees.",
            status_code=409,
        )
    row = dict(rows[0])
    if str(row["employee_code"]) != employee_code or str(row["security_user_id"]) != security_user_id:
        raise AttendanceRuleError(
            "EMPLOYEE_IMPORT_CONFLICT",
            "Employee code or Security user is already linked to another employee.",
            status_code=409,
        )
    return row


def preview_employee_workbook(
    connection: Connection,
    *,
    filename: str,
    data: bytes,
    actor_user_id: str,
) -> dict[str, object]:
    if not data:
        raise AttendanceRuleError("IMPORT_EMPTY", "Workbook is empty.", status_code=400)
    try:
        workbook = load_workbook(BytesIO(data), data_only=True)
    except Exception as exc:
        raise AttendanceRuleError(
            "IMPORT_INVALID",
            "Workbook could not be read.",
            status_code=400,
        ) from exc
    if "Employees" not in workbook.sheetnames:
        raise AttendanceRuleError(
            "IMPORT_SHEET_MISSING",
            "Workbook must contain an Employees sheet.",
            status_code=400,
        )
    sheet = workbook["Employees"]
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    if headers[: len(_HEADERS)] != _HEADERS:
        raise AttendanceRuleError(
            "IMPORT_HEADERS_INVALID",
            "Employee workbook columns do not match the current template.",
            status_code=400,
        )

    import_id = uuid4()
    planned: list[dict[str, object]] = []
    seen_codes: set[str] = set()
    seen_users: set[str] = set()

    for row_number, cells in enumerate(
        sheet.iter_rows(min_row=2, max_col=len(_HEADERS), values_only=True),
        start=2,
    ):
        if not any(value not in (None, "") for value in cells):
            continue
        payload, messages = _normalize_row(list(cells), row_number)
        action = "ERROR" if messages else "CREATE"
        employee_code = str(payload.get("employee_code") or "")
        security_user_id = str(payload.get("security_user_id") or "")
        if not messages:
            if employee_code in seen_codes:
                messages.append("EmployeeCode is duplicated in this workbook")
            if security_user_id in seen_users:
                messages.append("SecurityUserId is duplicated in this workbook")
            seen_codes.add(employee_code)
            seen_users.add(security_user_id)

        if not messages:
            try:
                existing = _existing_employee(
                    connection,
                    employee_code=employee_code,
                    security_user_id=security_user_id,
                )
                action = "UPDATE" if existing is not None else "CREATE"
                if existing is not None:
                    payload["employee_id"] = str(existing["employee_id"])
            except AttendanceRuleError as exc:
                messages.append(exc.detail)
                action = "ERROR"

        planned.append(
            {
                "row_number": row_number,
                "employee_code": employee_code or None,
                "action": "ERROR" if messages else action,
                "payload": payload,
                "messages": messages,
            }
        )

    if not planned:
        raise AttendanceRuleError(
            "IMPORT_NO_ROWS",
            "Workbook does not contain employee rows.",
            status_code=400,
        )

    connection.execute(
        text(
            """
            INSERT INTO verigence_attendance.bulk_imports (
                import_id,filename,status,uploaded_by_user_id
            ) VALUES (
                :import_id,:filename,'PREVIEW_READY',CAST(:actor AS uuid)
            )
            """
        ),
        {
            "import_id": import_id,
            "filename": filename[:255],
            "actor": actor_user_id,
        },
    )
    for item in planned:
        connection.execute(
            text(
                """
                INSERT INTO verigence_attendance.bulk_import_rows (
                    import_id,row_number,action,employee_code,payload_json,messages_json
                ) VALUES (
                    :import_id,:row_number,:action,:employee_code,
                    CAST(:payload AS jsonb),CAST(:messages AS jsonb)
                )
                """
            ),
            {
                "import_id": import_id,
                "row_number": item["row_number"],
                "action": item["action"],
                "employee_code": item["employee_code"],
                "payload": __import__("json").dumps(item["payload"], default=str),
                "messages": __import__("json").dumps(item["messages"]),
            },
        )
    return import_preview(connection, import_id)


def import_preview(connection: Connection, import_id: UUID) -> dict[str, object]:
    header = connection.execute(
        text(
            """
            SELECT import_id,filename,status,uploaded_at_utc,applied_at_utc
            FROM verigence_attendance.bulk_imports
            WHERE import_id=:import_id
            """
        ),
        {"import_id": import_id},
    ).mappings().first()
    if header is None:
        raise AttendanceNotFoundError("Employee import not found.")
    rows = [
        dict(row)
        for row in connection.execute(
            text(
                """
                SELECT row_number,action,employee_code,payload_json,messages_json,applied
                FROM verigence_attendance.bulk_import_rows
                WHERE import_id=:import_id
                ORDER BY row_number
                """
            ),
            {"import_id": import_id},
        ).mappings()
    ]
    counts = {"CREATE": 0, "UPDATE": 0, "UNCHANGED": 0, "ERROR": 0}
    for row in rows:
        counts[str(row["action"])] = counts.get(str(row["action"]), 0) + 1
    return {
        "importId": header["import_id"],
        "filename": header["filename"],
        "status": header["status"],
        "counts": counts,
        "rows": [
            {
                "rowNumber": row["row_number"],
                "employeeCode": row["employee_code"],
                "action": row["action"],
                "messages": row["messages_json"],
                "applied": row["applied"],
            }
            for row in rows
        ],
    }


def _apply_update(
    connection: Connection,
    *,
    payload: dict[str, object],
    actor_user_id: str,
) -> None:
    employee_id = UUID(str(payload["employee_id"]))
    work_location_id = payload.get("work_location_id")
    if work_location_id:
        valid_location = connection.execute(
            text(
                """
                SELECT 1 FROM verigence_attendance.work_locations
                WHERE location_id=CAST(:location_id AS uuid) AND status='ACTIVE'
                """
            ),
            {"location_id": work_location_id},
        ).scalar_one_or_none()
        if valid_location is None:
            raise AttendanceRuleError(
                "WORK_LOCATION_INVALID",
                "Selected work location is not active.",
                status_code=400,
            )

    connection.execute(
        text(
            """
            UPDATE verigence_attendance.employees
            SET display_name=:display_name,
                primary_email=:primary_email,
                mobile=:mobile,
                joining_date=:joining_date,
                tl_user_id=CAST(:tl_user_id AS uuid),
                pmo_user_id=CAST(:pmo_user_id AS uuid),
                project_tenant_id=CAST(:project_tenant_id AS uuid),
                work_location_id=CAST(:work_location_id AS uuid),
                bank_account_masked=:bank_account_masked,
                pan_masked=:pan_masked,
                aadhaar_masked=:aadhaar_masked,
                updated_at_utc=now()
            WHERE employee_id=:employee_id
            """
        ),
        {
            "employee_id": employee_id,
            "display_name": payload["display_name"],
            "primary_email": payload.get("primary_email"),
            "mobile": payload.get("mobile"),
            "joining_date": payload["joining_date"],
            "tl_user_id": payload.get("tl_user_id"),
            "pmo_user_id": payload.get("pmo_user_id"),
            "project_tenant_id": payload.get("project_tenant_id"),
            "work_location_id": payload.get("work_location_id"),
            "bank_account_masked": payload.get("bank_account_masked"),
            "pan_masked": payload.get("pan_masked"),
            "aadhaar_masked": payload.get("aadhaar_masked"),
        },
    )

    effective_from = date.fromisoformat(str(payload["salary_effective_from"]))
    salary = {
        "basic_salary": Decimal(str(payload["basic_salary"])),
        "hra": Decimal(str(payload["hra"])),
        "allowances": Decimal(str(payload["allowances"])),
        "other_earnings": Decimal(str(payload["other_earnings"])),
        "fixed_deductions": Decimal(str(payload["fixed_deductions"])),
    }
    current = connection.execute(
        text(
            """
            SELECT salary_structure_id,effective_from,basic_salary,hra,allowances,
                   other_earnings,fixed_deductions
            FROM verigence_attendance.salary_structures
            WHERE employee_id=:employee_id
              AND effective_from<=:effective_from
              AND (effective_to IS NULL OR effective_to>=:effective_from)
            ORDER BY effective_from DESC
            LIMIT 1
            FOR UPDATE
            """
        ),
        {"employee_id": employee_id, "effective_from": effective_from},
    ).mappings().first()
    if current is not None:
        same = all(Decimal(str(current[key])) == salary[key] for key in salary)
        if same:
            return
        if current["effective_from"] == effective_from:
            connection.execute(
                text(
                    """
                    UPDATE verigence_attendance.salary_structures
                    SET basic_salary=:basic_salary,hra=:hra,allowances=:allowances,
                        other_earnings=:other_earnings,fixed_deductions=:fixed_deductions
                    WHERE salary_structure_id=:salary_structure_id
                    """
                ),
                {**salary, "salary_structure_id": current["salary_structure_id"]},
            )
            return
        connection.execute(
            text(
                """
                UPDATE verigence_attendance.salary_structures
                SET effective_to=:effective_to
                WHERE salary_structure_id=:salary_structure_id
                """
            ),
            {
                "effective_to": effective_from - timedelta(days=1),
                "salary_structure_id": current["salary_structure_id"],
            },
        )
    connection.execute(
        text(
            """
            INSERT INTO verigence_attendance.salary_structures (
                employee_id,effective_from,basic_salary,hra,allowances,
                other_earnings,fixed_deductions,created_by_user_id
            ) VALUES (
                :employee_id,:effective_from,:basic_salary,:hra,:allowances,
                :other_earnings,:fixed_deductions,CAST(:actor AS uuid)
            )
            ON CONFLICT (employee_id,effective_from) DO UPDATE SET
                basic_salary=EXCLUDED.basic_salary,
                hra=EXCLUDED.hra,
                allowances=EXCLUDED.allowances,
                other_earnings=EXCLUDED.other_earnings,
                fixed_deductions=EXCLUDED.fixed_deductions
            """
        ),
        {
            "employee_id": employee_id,
            "effective_from": effective_from,
            **salary,
            "actor": actor_user_id,
        },
    )


def apply_employee_import(
    connection: Connection,
    *,
    import_id: UUID,
    actor_user_id: str,
) -> dict[str, object]:
    import_row = connection.execute(
        text(
            """
            SELECT import_id,status
            FROM verigence_attendance.bulk_imports
            WHERE import_id=:import_id
            FOR UPDATE
            """
        ),
        {"import_id": import_id},
    ).mappings().first()
    if import_row is None:
        raise AttendanceNotFoundError("Employee import not found.")
    if import_row["status"] == "APPLIED":
        return import_preview(connection, import_id)
    if import_row["status"] != "PREVIEW_READY":
        raise AttendanceRuleError(
            "IMPORT_STATE_INVALID",
            "Employee import is not ready to apply.",
        )

    rows = [
        dict(row)
        for row in connection.execute(
            text(
                """
                SELECT import_row_id,row_number,action,payload_json,messages_json
                FROM verigence_attendance.bulk_import_rows
                WHERE import_id=:import_id
                ORDER BY row_number
                FOR UPDATE
                """
            ),
            {"import_id": import_id},
        ).mappings()
    ]
    if any(row["action"] == "ERROR" for row in rows):
        raise AttendanceRuleError(
            "IMPORT_HAS_ERRORS",
            "Fix workbook errors and upload it again before applying.",
            status_code=409,
        )

    for row in rows:
        payload = dict(row["payload_json"])
        if row["action"] == "CREATE":
            create_employee(
                connection,
                user_id=UUID(str(payload["security_user_id"])),
                employee_code=str(payload["employee_code"]),
                display_name=str(payload["display_name"]),
                primary_email=payload.get("primary_email"),
                mobile=payload.get("mobile"),
                joining_date=date.fromisoformat(str(payload["joining_date"])),
                tl_user_id=(
                    UUID(str(payload["tl_user_id"])) if payload.get("tl_user_id") else None
                ),
                pmo_user_id=(
                    UUID(str(payload["pmo_user_id"])) if payload.get("pmo_user_id") else None
                ),
                project_tenant_id=(
                    UUID(str(payload["project_tenant_id"]))
                    if payload.get("project_tenant_id")
                    else None
                ),
                work_location_id=(
                    UUID(str(payload["work_location_id"]))
                    if payload.get("work_location_id")
                    else None
                ),
                bank_account_masked=payload.get("bank_account_masked"),
                pan_masked=payload.get("pan_masked"),
                aadhaar_masked=payload.get("aadhaar_masked"),
                salary={
                    "basic_salary": Decimal(str(payload["basic_salary"])),
                    "hra": Decimal(str(payload["hra"])),
                    "allowances": Decimal(str(payload["allowances"])),
                    "other_earnings": Decimal(str(payload["other_earnings"])),
                    "fixed_deductions": Decimal(str(payload["fixed_deductions"])),
                },
                actor_user_id=actor_user_id,
            )
        elif row["action"] == "UPDATE":
            _apply_update(
                connection,
                payload=payload,
                actor_user_id=actor_user_id,
            )
        connection.execute(
            text(
                """
                UPDATE verigence_attendance.bulk_import_rows
                SET applied=true
                WHERE import_row_id=:import_row_id
                """
            ),
            {"import_row_id": row["import_row_id"]},
        )

    connection.execute(
        text(
            """
            UPDATE verigence_attendance.bulk_imports
            SET status='APPLIED',applied_at_utc=now()
            WHERE import_id=:import_id
            """
        ),
        {"import_id": import_id},
    )
    return import_preview(connection, import_id)
