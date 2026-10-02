from __future__ import annotations

from datetime import date
from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import Connection, text

from verigence_security.attendance.employee.errors import AttendanceNotFoundError


def _format_sheet(worksheet: Worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    for index, column in enumerate(worksheet.iter_cols(), start=1):
        width = max(
            (len(str(cell.value)) if cell.value is not None else 0)
            for cell in column
        )
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            max(width + 2, 10),
            36,
        )


def _workbook_bytes(workbook: Workbook) -> bytes:
    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


def attendance_report(
    connection: Connection,
    *,
    start_date: date,
    end_date: date,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"
    sheet.append(
        [
            "Employee Code",
            "Employee Name",
            "Date",
            "Status",
            "Present Fraction",
            "Check In",
            "Check Out",
            "HR Review Status",
            "HR Review Comment",
        ]
    )
    rows = connection.execute(
        text(
            """
            SELECT e.employee_code,e.display_name,a.attendance_date,a.status,
                   a.present_fraction,a.check_in_at_utc,a.check_out_at_utc,
                   a.hr_review_status,a.hr_review_comment
            FROM verigence_attendance.employees e
            LEFT JOIN verigence_attendance.attendance_days a
              ON a.employee_id=e.employee_id
             AND a.attendance_date BETWEEN :start_date AND :end_date
            WHERE e.employment_status='ACTIVE'
            ORDER BY e.employee_code,a.attendance_date
            """
        ),
        {"start_date": start_date, "end_date": end_date},
    ).mappings()
    for row in rows:
        sheet.append(
            [
                row["employee_code"],
                row["display_name"],
                row["attendance_date"],
                row["status"],
                float(row["present_fraction"]) if row["present_fraction"] is not None else None,
                row["check_in_at_utc"],
                row["check_out_at_utc"],
                row["hr_review_status"],
                row["hr_review_comment"],
            ]
        )
    _format_sheet(sheet)
    return _workbook_bytes(workbook)


def payroll_report(connection: Connection, *, run_id: UUID) -> bytes:
    run = connection.execute(
        text(
            """
            SELECT payroll_run_id,payroll_month,status
            FROM verigence_attendance.payroll_runs
            WHERE payroll_run_id=:run_id
            """
        ),
        {"run_id": run_id},
    ).mappings().first()
    if run is None:
        raise AttendanceNotFoundError("Payroll run not found.")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Payroll"
    sheet.append(
        [
            "Employee Code",
            "Employee Name",
            "Payroll Month",
            "Status",
            "Scheduled Days",
            "Present Days",
            "Paid Leave",
            "Unpaid Leave",
            "Payable Days",
            "Basic",
            "HRA",
            "Allowances",
            "Other Earnings",
            "LOP Amount",
            "Gross Amount",
            "Employee PF",
            "Employee ESI",
            "Professional Tax",
            "TDS",
            "Other Deductions",
            "Total Deductions",
            "Net Amount",
            "Employer PF",
            "Employer EPS",
            "Employer ESI",
            "Gratuity Provision",
            "Employer Cost",
        ]
    )
    rows = connection.execute(
        text(
            """
            SELECT e.employee_code,e.display_name,i.scheduled_days,i.present_days,
                   i.paid_leave_days,i.unpaid_leave_days,i.payable_days,
                   i.basic_amount,i.hra_amount,i.allowances_amount,
                   i.other_earnings_amount,i.lop_amount,i.gross_amount,
                   i.employee_pf,i.employee_esi,i.professional_tax,i.tds_amount,
                   i.other_deductions,i.deduction_amount,i.net_amount,
                   i.employer_pf,i.employer_eps,i.employer_esi,
                   i.gratuity_provision,i.employer_cost
            FROM verigence_attendance.payroll_items i
            JOIN verigence_attendance.employees e ON e.employee_id=i.employee_id
            WHERE i.payroll_run_id=:run_id
            ORDER BY e.employee_code
            """
        ),
        {"run_id": run_id},
    ).mappings()
    for row in rows:
        sheet.append(
            [
                row["employee_code"],
                row["display_name"],
                run["payroll_month"],
                run["status"],
                float(row["scheduled_days"]),
                float(row["present_days"]),
                float(row["paid_leave_days"]),
                float(row["unpaid_leave_days"]),
                float(row["payable_days"]),
                float(row["basic_amount"]),
                float(row["hra_amount"]),
                float(row["allowances_amount"]),
                float(row["other_earnings_amount"]),
                float(row["lop_amount"]),
                float(row["gross_amount"]),
                float(row["employee_pf"]),
                float(row["employee_esi"]),
                float(row["professional_tax"]),
                float(row["tds_amount"]),
                float(row["other_deductions"]),
                float(row["deduction_amount"]),
                float(row["net_amount"]),
                float(row["employer_pf"]),
                float(row["employer_eps"]),
                float(row["employer_esi"]),
                float(row["gratuity_provision"]),
                float(row["employer_cost"]),
            ]
        )
    for row in sheet.iter_rows(min_row=2, min_col=10, max_col=27):
        for cell in row:
            cell.number_format = '#,##0.00'
    _format_sheet(sheet)
    return _workbook_bytes(workbook)
